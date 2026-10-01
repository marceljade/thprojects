"""Vorlagen, Bearbeiter, Aufgabenarten, Feiertage, Einstellungen, Export, Backup."""
from __future__ import annotations

import csv
import io
import json
import shutil
from datetime import date, datetime

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from .. import config, models, scheduling
from ..enums import (DUE_STATE_LABELS, PRIORITY_LABELS, PROJECT_CATEGORY_LABELS, PROJECT_STATUS_LABELS,
                     TASK_STATUS_LABELS, Priority, ProjectCategory, ProjectStatus, TaskStatus)
from ..schemas import (BackupOut, HolidayIn, HolidayOut, SettingsIn, SettingsOut, TaskTypeIn, TaskTypeOut,
                       TemplateIn, TemplateOut, TemplatePreviewRow, TemplateTaskOut, UserIn, UserOut)
from . import common


# ---------------------------------------------------------------- Vorlagen --
def template_out(t: models.ProjectTemplate) -> TemplateOut:
    return TemplateOut(id=t.id, name=t.name, description=t.description or "", category=t.category, sort_order=t.sort_order,
                       tasks=[TemplateTaskOut(id=tt.id, sort_order=tt.sort_order, title=tt.title, task_type_id=tt.task_type_id,
                                              task_type_name=tt.task_type.name if tt.task_type else None,
                                              offset_workdays=tt.offset_workdays, weight=tt.weight,
                                              assignee_role=tt.assignee_role, depends_on_previous=tt.depends_on_previous)
                              for tt in sorted(t.tasks, key=lambda x: x.sort_order)])


def list_templates(db: Session) -> list[TemplateOut]:
    rows = db.scalars(select(models.ProjectTemplate).options(selectinload(models.ProjectTemplate.tasks)).order_by(models.ProjectTemplate.sort_order, models.ProjectTemplate.name)).unique().all()
    return [template_out(t) for t in rows]


def save_template(db: Session, data: TemplateIn, template_id: int | None = None) -> TemplateOut:
    if template_id is None:
        t = models.ProjectTemplate(name=data.name, description=data.description, category=data.category, sort_order=data.sort_order)
        db.add(t)
    else:
        t = db.get(models.ProjectTemplate, template_id)
        if t is None:
            raise HTTPException(404, "Vorlage nicht gefunden.")
        t.name, t.description, t.category, t.sort_order = data.name, data.description, data.category, data.sort_order
        t.tasks.clear()
    for i, tt in enumerate(data.tasks, start=1):
        t.tasks.append(models.TemplateTask(sort_order=i, title=tt.title, task_type_id=tt.task_type_id, offset_workdays=tt.offset_workdays,
                                           weight=tt.weight, assignee_role=tt.assignee_role or "PL", depends_on_previous=tt.depends_on_previous))
    db.commit()
    db.refresh(t)
    return template_out(t)


def delete_template(db: Session, template_id: int) -> None:
    t = db.get(models.ProjectTemplate, template_id)
    if t is None:
        raise HTTPException(404, "Vorlage nicht gefunden.")
    db.delete(t)
    db.commit()


def preview_template(db: Session, template_id: int, deadline: date | None) -> list[TemplatePreviewRow]:
    t = db.get(models.ProjectTemplate, template_id)
    if t is None:
        raise HTTPException(404, "Vorlage nicht gefunden.")
    p = common.params(db)
    tasks = sorted(t.tasks, key=lambda x: x.sort_order)
    dues = scheduling.template_due_dates(deadline, [tt.offset_workdays for tt in tasks], p.holidays, p.today)
    return [TemplatePreviewRow(title=tt.title, task_type_name=tt.task_type.name if tt.task_type else None,
                               offset_workdays=tt.offset_workdays, due_date=d,
                               weight=tt.weight if tt.weight is not None else (tt.task_type.default_weight if tt.task_type else 1.0),
                               assignee_role=tt.assignee_role) for tt, d in zip(tasks, dues)]


# -------------------------------------------------------------- Bearbeiter --
def list_users(db: Session) -> list[UserOut]:
    return [UserOut.model_validate(u) for u in db.scalars(select(models.User).order_by(models.User.sort_order, models.User.code)).all()]


def save_user(db: Session, data: UserIn, user_id: int | None = None) -> UserOut:
    existing = db.scalar(select(models.User).where(models.User.code == data.code))
    if existing is not None and existing.id != user_id:
        raise HTTPException(409, f"Das Kürzel {data.code} ist bereits vergeben.")
    if user_id is None:
        u = models.User(**data.model_dump())
        db.add(u)
    else:
        u = db.get(models.User, user_id)
        if u is None:
            raise HTTPException(404, "Bearbeiter nicht gefunden.")
        for k, v in data.model_dump().items():
            setattr(u, k, v)
    db.commit()
    db.refresh(u)
    return UserOut.model_validate(u)


def delete_user(db: Session, user_id: int) -> None:
    u = db.get(models.User, user_id)
    if u is None:
        raise HTTPException(404, "Bearbeiter nicht gefunden.")
    used = db.scalar(select(models.Task).where(models.Task.assignee_id == user_id).limit(1)) or db.scalar(select(models.Project).where(models.Project.assignee_id == user_id).limit(1))
    if used:
        raise HTTPException(409, "Dieser Bearbeiter ist noch Projekten oder Aufgaben zugeordnet. Stattdessen auf inaktiv setzen.")
    db.delete(u)
    db.commit()


# ------------------------------------------------------------ Aufgabenarten --
def list_task_types(db: Session) -> list[TaskTypeOut]:
    return [TaskTypeOut.model_validate(t) for t in db.scalars(select(models.TaskType).order_by(models.TaskType.sort_order, models.TaskType.name)).all()]


def save_task_type(db: Session, data: TaskTypeIn, type_id: int | None = None) -> TaskTypeOut:
    existing = db.scalar(select(models.TaskType).where(models.TaskType.name == data.name))
    if existing is not None and existing.id != type_id:
        raise HTTPException(409, f"Die Aufgabenart {data.name} gibt es schon.")
    if type_id is None:
        t = models.TaskType(**data.model_dump())
        db.add(t)
    else:
        t = db.get(models.TaskType, type_id)
        if t is None:
            raise HTTPException(404, "Aufgabenart nicht gefunden.")
        for k, v in data.model_dump().items():
            setattr(t, k, v)
    db.commit()
    db.refresh(t)
    return TaskTypeOut.model_validate(t)


def delete_task_type(db: Session, type_id: int) -> None:
    t = db.get(models.TaskType, type_id)
    if t is None:
        raise HTTPException(404, "Aufgabenart nicht gefunden.")
    db.delete(t)
    db.commit()


# ---------------------------------------------------------------- Feiertage --
def list_holidays(db: Session) -> list[HolidayOut]:
    return [HolidayOut.model_validate(h) for h in db.scalars(select(models.Holiday).order_by(models.Holiday.date)).all()]


def add_holiday(db: Session, data: HolidayIn) -> HolidayOut:
    if db.scalar(select(models.Holiday).where(models.Holiday.date == data.date)):
        raise HTTPException(409, "Für dieses Datum gibt es schon einen Eintrag.")
    h = models.Holiday(date=data.date, name=data.name)
    db.add(h)
    db.commit()
    db.refresh(h)
    return HolidayOut.model_validate(h)


def delete_holiday(db: Session, holiday_id: int) -> None:
    h = db.get(models.Holiday, holiday_id)
    if h is None:
        raise HTTPException(404, "Feiertag nicht gefunden.")
    db.delete(h)
    db.commit()


# ------------------------------------------------------------- Einstellungen --
def update_settings(db: Session, data: SettingsIn) -> SettingsOut:
    if data.my_user_id is not None:
        if db.get(models.User, data.my_user_id) is None:
            raise HTTPException(422, "Bearbeiter nicht gefunden.")
        common.set_setting(db, "my_user_id", str(data.my_user_id))
    if data.base_path is not None:
        common.set_setting(db, "base_path", data.base_path.strip())
    if data.warn_workdays is not None:
        common.set_setting(db, "warn_workdays", str(data.warn_workdays))
    if data.warn_project_workdays is not None:
        common.set_setting(db, "warn_project_workdays", str(data.warn_project_workdays))
    if data.auto_backup is not None:
        common.set_setting(db, "auto_backup", "1" if data.auto_backup else "0")
    if data.dashboard_widgets is not None:
        common.set_setting(db, "dashboard_widgets", json.dumps(data.dashboard_widgets))
    db.commit()
    return common.settings_out(db)


def meta_lists() -> dict:
    return {
        "project_categories": [{"value": k.value, "label": v} for k, v in PROJECT_CATEGORY_LABELS.items()],
        "project_statuses": [{"value": k.value, "label": v, "active": k in scheduling.ACTIVE_PROJECT_STATUSES} for k, v in PROJECT_STATUS_LABELS.items()],
        "task_statuses": [{"value": k.value, "label": v, "open": k in scheduling.OPEN_TASK_STATUSES} for k, v in TASK_STATUS_LABELS.items()],
        "priorities": [{"value": k.value, "label": v} for k, v in PRIORITY_LABELS.items()],
        "due_states": [{"value": k.value, "label": v} for k, v in DUE_STATE_LABELS.items()],
        "signals": [{"value": "rot", "label": "Kritisch"}, {"value": "gelb", "label": "Frist nähert sich"}, {"value": "gruen", "label": "Im Plan"},
                    {"value": "grau", "label": "Nicht aktiv"}, {"value": "fertig", "label": "Abgeschlossen"}],
    }


# ------------------------------------------------------------------ Export --
def export_excel(db: Session) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    p = common.params(db)
    wb = Workbook()

    def sheet(title: str, headers: list[str], rows: list[list], name: str):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True)
        for i, h in enumerate(headers, start=1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None][:200])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
        if rows:
            ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
            tbl = Table(displayName=name, ref=ref)
            tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            ws.add_table(tbl)
        ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, (date, datetime)):
                    c.number_format = "DD.MM.YYYY" if isinstance(c.value, date) and not isinstance(c.value, datetime) else "DD.MM.YYYY HH:MM"

    prows = []
    for pr in common.load_projects(db):
        o = common.project_out(pr, p)
        prows.append([o.project_number, o.name, PROJECT_CATEGORY_LABELS[ProjectCategory(o.category)], o.client, o.assignee_code, o.participants,
                      PROJECT_STATUS_LABELS[ProjectStatus(o.status)], PRIORITY_LABELS[Priority(o.priority)], o.request_date, o.offer_date,
                      o.order_date, o.offered_weeks, o.contractual_deadline, o.target_deadline, o.deadline, o.completed_at,
                      o.progress, o.task_count, o.open_count, o.overdue_count, o.next_task_title, o.next_due,
                      {"rot": "Rot", "gelb": "Gelb", "gruen": "Grün", "grau": "Nicht aktiv", "fertig": "Abgeschlossen"}[o.signal.value],
                      " · ".join(o.reasons + o.warnings), o.folder_path, o.remarks])
    sheet("Projekte", ["Projektnr", "Projektname", "Kategorie", "Auftraggeber", "Projektleiter", "Beteiligte", "Status", "Priorität",
                       "Anfrage", "Angebot", "Auftrag", "Dauer (Wochen)", "Fertigstellung vertraglich", "Fertigstellung Ziel", "Fertigstellung",
                       "Fertigstellung tatsächlich", "Fortschritt %", "Aufgaben", "Offen", "Überfällig", "Nächste Aufgabe", "Nächste Frist",
                       "Ampel", "Hinweise", "Projektordner", "Bemerkung"], prows, "tblProjekte")

    trows = []
    stmt = select(models.Task).options(selectinload(models.Task.project).selectinload(models.Project.tasks)).order_by(models.Task.project_id, models.Task.sort_order)
    for t in db.scalars(stmt).unique().all():
        o = common.task_out(t, p)
        trows.append([o.id, o.project_number, o.project_name, o.title, o.task_type_name, o.assignee_code, TASK_STATUS_LABELS[TaskStatus(o.status)],
                      PRIORITY_LABELS[Priority(o.priority)], o.start_date, o.due_date, o.kw, o.completed_at, o.progress, o.weight,
                      o.predecessor_id, DUE_STATE_LABELS[o.due_state], o.waiting_for, o.waiting_on, o.waiting_since, o.description])
    sheet("Aufgaben", ["ID", "Projektnr", "Projektname", "Aufgabe", "Aufgabenart", "Verantwortlich", "Status", "Priorität", "Start", "Frist", "KW",
                       "Erledigt am", "Fortschritt %", "Gewicht", "Vorgänger", "Fälligkeit", "Wartet auf", "Von", "Seit", "Beschreibung"], trows, "tblAufgaben")

    nrows = [[n.id, n.project.project_number if n.project else "", n.task.title if n.task else "", n.author.code if n.author else "", n.created_at, n.content]
             for n in db.scalars(select(models.Note).order_by(models.Note.created_at)).all()]
    sheet("Notizen", ["ID", "Projektnr", "Aufgabe", "Autor", "Datum", "Text"], nrows, "tblNotizen")

    arows = [[a.created_at, a.user.code if a.user else "", a.project.project_number if a.project else "", a.task.title if a.task else "",
              a.action, a.field, a.old_value, a.new_value, a.details]
             for a in db.scalars(select(models.ActivityLog).order_by(models.ActivityLog.created_at)).all()]
    sheet("Aktivitäten", ["Zeitpunkt", "Benutzer", "Projektnr", "Aufgabe", "Aktion", "Feld", "Alt", "Neu", "Details"], arows, "tblAktivitaeten")

    wb.remove(wb["Sheet"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_csv(db: Session, what: str) -> str:
    p = common.params(db)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\n")
    if what == "projects":
        w.writerow(["Projektnr", "Projektname", "Kategorie", "Auftraggeber", "Projektleiter", "Status", "Priorität", "Anfrage", "Angebot", "Auftrag",
                    "Dauer (Wochen)", "Fertigstellung", "Fortschritt %", "Nächste Aufgabe", "Nächste Frist", "Ampel", "Projektordner"])
        for pr in common.load_projects(db):
            o = common.project_out(pr, p)
            w.writerow([o.project_number, o.name, o.category.value, o.client, o.assignee_code or "", o.status.value, o.priority.value,
                        _d(o.request_date), _d(o.offer_date), _d(o.order_date), o.offered_weeks or "", _d(o.deadline), o.progress,
                        o.next_task_title or "", _d(o.next_due), o.signal.value, o.folder_path])
    else:
        w.writerow(["ID", "Projektnr", "Aufgabe", "Aufgabenart", "Verantwortlich", "Status", "Priorität", "Start", "Frist", "KW", "Erledigt am", "Fortschritt %", "Gewicht"])
        stmt = select(models.Task).options(selectinload(models.Task.project).selectinload(models.Project.tasks)).order_by(models.Task.project_id, models.Task.sort_order)
        for t in db.scalars(stmt).unique().all():
            o = common.task_out(t, p)
            w.writerow([o.id, o.project_number, o.title, o.task_type_name or "", o.assignee_code or "", o.status.value, o.priority.value,
                        _d(o.start_date), _d(o.due_date), o.kw or "", _d(o.completed_at), o.progress, o.weight])
    return buf.getvalue()


def _d(d: date | None) -> str:
    return d.strftime("%d.%m.%Y") if d else ""


# ------------------------------------------------------------------ Backup --
def create_backup(reason: str = "manuell") -> BackupOut:
    config.ensure_dirs()
    if not config.DB_PATH.exists():
        raise HTTPException(404, "Es gibt noch keine Datenbankdatei.")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    target = config.BACKUP_DIR / f"projekte_{stamp}.db"
    # sqlite3-Backup-API, damit auch bei WAL eine konsistente Kopie entsteht
    import sqlite3
    src = sqlite3.connect(str(config.DB_PATH))
    dst = sqlite3.connect(str(target))
    with dst:
        src.backup(dst)
    dst.close()
    src.close()
    backups = sorted(config.BACKUP_DIR.glob("projekte_*.db"))
    for old in backups[:-config.BACKUPS_KEEP]:
        old.unlink(missing_ok=True)
    st = target.stat()
    return BackupOut(file=target.name, size_bytes=st.st_size, created_at=datetime.fromtimestamp(st.st_mtime))


def list_backups() -> list[BackupOut]:
    config.ensure_dirs()
    out = []
    for f in sorted(config.BACKUP_DIR.glob("projekte_*.db"), reverse=True):
        st = f.stat()
        out.append(BackupOut(file=f.name, size_bytes=st.st_size, created_at=datetime.fromtimestamp(st.st_mtime)))
    return out


def daily_backup_if_due(db: Session) -> None:
    if common.get_setting(db, "auto_backup", "1") != "1":
        return
    last = common.get_setting(db, "last_backup", "")
    today = date.today().isoformat()
    if last == today or not config.DB_PATH.exists():
        return
    try:
        create_backup("automatisch")
        common.set_setting(db, "last_backup", today)
        db.commit()
    except Exception:  # noqa: BLE001 – Backup darf den Start nie verhindern
        db.rollback()
